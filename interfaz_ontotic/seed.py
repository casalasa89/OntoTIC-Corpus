"""Crea la base local de los 43 actores y el corpus OntoTIC."""
from __future__ import annotations
import csv
import json
import sqlite3
from pathlib import Path
from openpyxl import load_workbook

BASE=Path(__file__).resolve().parent
SOURCE=BASE/'sources'
DB=BASE/'data'/'ontotic.sqlite'
CAPS=('Investigación','Desarrollo','Mercadeo','Difusión y transferencia')

def build_database(db=DB):
    db.parent.mkdir(parents=True,exist_ok=True)
    if db.exists():db.unlink()
    cx=sqlite3.connect(db)
    cx.executescript('''
    PRAGMA foreign_keys=ON;
    CREATE TABLE actors(id INTEGER PRIMARY KEY,name TEXT NOT NULL UNIQUE,role TEXT NOT NULL,
       origin TEXT NOT NULL DEFAULT 'estudio',created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
    CREATE TABLE baseline(actor_id INTEGER PRIMARY KEY REFERENCES actors(id),lda_docs INTEGER NOT NULL,
       signaled_docs INTEGER NOT NULL,raw_json TEXT NOT NULL,adjusted_json TEXT NOT NULL);
    CREATE TABLE documents(id INTEGER PRIMARY KEY,source_id TEXT UNIQUE,actor_id INTEGER NOT NULL REFERENCES actors(id),
       title TEXT NOT NULL,publication_year INTEGER,publication_date TEXT,source_name TEXT,url TEXT,
       text_body TEXT,extracted INTEGER NOT NULL DEFAULT 0,analytical_status TEXT NOT NULL,
       signal_json TEXT,evidence_json TEXT,created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
    CREATE INDEX doc_actor ON documents(actor_id);
    CREATE TABLE dyads(actor_a INTEGER NOT NULL REFERENCES actors(id),actor_b INTEGER NOT NULL REFERENCES actors(id),
       potential REAL,reliability REAL,priority REAL,structural_rank INTEGER,evidence_rank INTEGER,
       PRIMARY KEY(actor_a,actor_b));
    CREATE INDEX dyad_b ON dyads(actor_b);
    CREATE TABLE metadata(key TEXT PRIMARY KEY,value TEXT NOT NULL);
    ''')
    matrix=load_workbook(SOURCE/'Matriz_Diadica_OntoTIC_43_actores_2010_2026_actualizada.xlsx',read_only=True,data_only=True)
    role={}
    for r in matrix['Díadas 903'].values:
        if r[2] and r[4] and r[2]!='Actor A':role[r[2]]=r[3];role[r[4]]=r[5]
    with (SOURCE/'perfiles_actores.csv').open(encoding='utf-8-sig',newline='') as stream:
        profiles=list(csv.DictReader(stream))
    actor_id={}
    for r in profiles:
        name=r['actor'];cur=cx.execute('INSERT INTO actors(name,role) VALUES(?,?)',(name,role.get(name,'Sin clasificar')))
        actor_id[name]=cur.lastrowid
        raw=[float(r['bruto_'+c] or 0) for c in CAPS]
        adjusted=[float(r['ajustado_'+c] or 0) for c in CAPS]
        cx.execute('INSERT INTO baseline VALUES(?,?,?,?,?)',(cur.lastrowid,int(r['documentos_lda']),int(r['documentos_con_señal']),json.dumps(raw),json.dumps(adjusted)))
    corpus=load_workbook(SOURCE/'Corpus_OntoTIC_2010_2026_actualizado_fuentes_institucionales.xlsx',read_only=True,data_only=True)
    it=iter(corpus['Corpus maestro 2010-2026'].values);header=next(it)
    batch=[]
    for row in it:
        if not row[0]:continue
        r=dict(zip(header,row));a=actor_id.get(r['actor'])
        if a is None:raise ValueError('Actor fuera de los 43: '+str(r['actor']))
        year=int(r['año']) if isinstance(r['año'],(int,float)) and r['año'] else None
        batch.append((r['documento_id'],a,str(r['título'] or ''),year,None,str(r['fuente'] or ''),str(r['url'] or ''),None,0,'Fuente histórica; no recalculada en la interfaz',None,None))
    cx.executemany('''INSERT INTO documents(source_id,actor_id,title,publication_year,publication_date,
        source_name,url,text_body,extracted,analytical_status,signal_json,evidence_json)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?)''',batch)
    add_dyads(cx,matrix,actor_id)
    prior=[0.0]*4;n=0
    with (SOURCE/'señales_documento.csv').open(encoding='utf-8-sig',newline='') as stream:
        for row in csv.DictReader(stream):
            if float(row['puntaje_lexico'] or 0)>0:
                n+=1
                for i,c in enumerate(CAPS):prior[i]+=float(row[c] or 0)
    prior=[x/n for x in prior]
    for key,val in {'prior':json.dumps(prior),'baseline_lda_docs':'8556','corpus_records':str(len(batch)),
                    'study_cutoff':'2026-08-31','author':'PhD. Carlos Andrés Salazar',
                    'method':'Perfil base LDA v6.13; documentos nuevos: señales léxicas y contracción n/(n+10)'}.items():
        cx.execute('INSERT INTO metadata VALUES(?,?)',(key,val))
    cx.commit();cx.close()
    return {'actors':len(actor_id),'documents':len(batch),'prior':prior}

def add_dyads(cx,matrix=None,actor_id=None):
    """Importa las 903 parejas del corte original, sin inferir relaciones observadas."""
    if matrix is None:matrix=load_workbook(SOURCE/'Matriz_Diadica_OntoTIC_43_actores_2010_2026_actualizada.xlsx',read_only=True,data_only=True)
    if actor_id is None:actor_id={r['name']:r['id'] for r in cx.execute('SELECT id,name FROM actors')}
    cx.execute('''CREATE TABLE IF NOT EXISTS dyads(actor_a INTEGER NOT NULL REFERENCES actors(id),actor_b INTEGER NOT NULL REFERENCES actors(id),
       potential REAL,reliability REAL,priority REAL,structural_rank INTEGER,evidence_rank INTEGER,
       PRIMARY KEY(actor_a,actor_b))''')
    cx.execute('CREATE INDEX IF NOT EXISTS dyad_b ON dyads(actor_b)')
    rows=iter(matrix['Díadas 903'].values);header=next(rows)
    payload=[]
    for row in rows:
        if not row[2] or not row[4]:continue
        r=dict(zip(header,row));a,b=actor_id[r['Actor A']],actor_id[r['Actor B']]
        payload.append((min(a,b),max(a,b),r['Complementariedad potencial'],r['Fiabilidad'],
                        r['Prioridad con evidencia'],r['Rango estructural'],r['Rango evidencia']))
    if len(payload)!=903 or sum(r[2] is not None for r in payload)!=741:raise ValueError('Matriz diádica inconsistente')
    cx.executemany('INSERT OR REPLACE INTO dyads VALUES(?,?,?,?,?,?,?)',payload)
    cx.commit()

if __name__=='__main__':print(build_database())
