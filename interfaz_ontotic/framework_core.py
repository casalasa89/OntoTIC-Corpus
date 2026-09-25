"""Análisis reproducible OntoTIC V6.13.

Cadena: corpus completo apto -> selección y estabilidad LDA -> alineación
léxica OntoTIC -> perfiles documentales de capacidad por actor -> similitud.
Jaccard representa solapamiento semántico/similitud, nunca complementariedad.
"""
from __future__ import annotations
import csv, itertools, json, math, re, sys, unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linear_sum_assignment
from scipy.sparse import csc_matrix
from sklearn.decomposition import LatentDirichletAllocation
from sklearn.feature_extraction.text import CountVectorizer, ENGLISH_STOP_WORDS
from sklearn.metrics.pairwise import cosine_similarity

SEEDS=(42,52,62)
KS=(8,10,12,15)
TOP_N=20
CORE=("Investigación","Desarrollo","Mercadeo","Difusión y transferencia")
CODE_TO_CORE={
 "RC":"Investigación","LC":"Investigación",
 "DC":"Desarrollo","MC":"Desarrollo","DM":"Desarrollo",
 "MK":"Mercadeo",
 "TC":"Difusión y transferencia","DS":"Difusión y transferencia","NC":"Difusión y transferencia",
}
ENABLERS={"OC":"Organizacional","RA":"Asignación de recursos","SP":"Planificación estratégica"}
CODE_ORDER=("RC","DC","MC","MK","OC","RA","SP","LC","TC","NC","DS","DM")
GENERIC={"modelo","proceso","sistema","gestion","desarrollo","capacidad","innovacion","aplicacion",
         "producto","servicio","tecnologia","technology","research","development","process","system","model"}
CONTEXT={"innovacion","tecnologia","tecnologico","i d","investigacion","producto","proceso","patente",
         "prototipo","mercado","transferencia","conocimiento","innovation","technology","research",
         "product","patent","prototype","market","transfer","knowledge"}
ES=set("""el la los las un una unos unas de del a al y o en por para con sin sobre entre como que se su sus
es son fue fueron ser ha han desde hacia durante mediante este esta estos estas lo le les mas ya pero si no
cada cual cuales e u esto ello aquel otras otros donde cuando mismo misma puede pueden presente estudio
trabajo resultados objetivo analisis evaluacion produccion colombiano colombia universidad tesis articulo
resumen documento datos uso utilizando mediante partir obtencion caracterizacion efectos efecto caso casos
diferentes estudios propuesta aplicado aplicada debido mayor tres medio valle cauca""".split())

def norm(s):
    s=unicodedata.normalize("NFKD",str(s or "").lower())
    s="".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+"," ",re.sub(r"[^a-z0-9+ ]+"," ",s)).strip()

def rows_after_header(ws,key):
    it=ws.iter_rows(values_only=True)
    for row in it:
        if row and row[0]==key:
            header=[str(x or "") for x in row];break
    else: raise ValueError(f"No se encontró encabezado {key}")
    for row in it:
        if row and row[0]: yield dict(zip(header,row))

def load_corpus(path):
    wb=load_workbook(path,read_only=True,data_only=True)
    actors=[str(r["actor"]).strip() for r in rows_after_header(wb["Cobertura 43"],"actor")]
    docs=[]
    for r in rows_after_header(wb["Corpus LDA"],"documento_id"):
        text=norm(f'{r.get("título","")} {r.get("texto_modelado","")}')
        if len(text)>=200:
            docs.append({"id":str(r["documento_id"]),"actor":str(r["actor"]).strip(),
                         "title":str(r.get("título") or ""),"year":r.get("año"),"text":text})
    return actors,docs

def load_lexicon(path):
    wb=load_workbook(path,read_only=True,data_only=True); ws=wb["Diccionario LDA"]
    head=None; data=[]
    for row in ws.iter_rows(values_only=True):
        if row and row[0]=="ID léxico": head=[str(x or "") for x in row];continue
        if head and row and row[0]: data.append(dict(zip(head,row)))
    forms=defaultdict(dict)
    for r in data:
        form=norm(r["Forma normalizada"] or r["Forma léxica"]); code=str(r["Código capacidad"])
        if not form:continue
        weight=float(r["Peso inicial"] or 1)
        forms[form][code]=max(weight,forms[form].get(code,0))
    return data,forms

def vectorize(texts):
    stops={norm(x) for x in ENGLISH_STOP_WORDS}|ES
    v=CountVectorizer(stop_words=sorted(stops),min_df=8,max_df=.72,max_features=6000,
       ngram_range=(1,2),token_pattern=r"(?u)\b[a-z][a-z0-9]+\b",strip_accents="unicode")
    return v,v.fit_transform(texts)

def topic_diversity(model,n=TOP_N):
    top=[set(np.argsort(w)[::-1][:n]) for w in model.components_]
    return len(set().union(*top))/(len(top)*n)

def npmi_coherence(model,X,n=10):
    B=csc_matrix(X.astype(bool)); N=X.shape[0]; vals=[]
    for weights in model.components_:
        ids=np.argsort(weights)[::-1][:n]
        for i,j in itertools.combinations(ids,2):
            di=B[:,i].nnz; dj=B[:,j].nnz
            both=B[:,i].multiply(B[:,j]).nnz
            if not both: vals.append(-1.0);continue
            pi,pj,pij=di/N,dj/N,both/N
            vals.append(math.log(pij/(pi*pj))/-math.log(pij))
    return float(np.mean(vals)) if vals else -1.0

def stability(models):
    sims=[]
    for a,b in itertools.combinations(models,2):
        C=cosine_similarity(a.components_,b.components_)
        ri,ci=linear_sum_assignment(-C);sims.append(float(C[ri,ci].mean()))
    return float(np.mean(sims)) if sims else 1.0

def fit_select(X):
    all_models={};metrics=[]
    for k in KS:
        ms=[]; per=[]; coh=[]; div=[]
        for seed in SEEDS:
            m=LatentDirichletAllocation(n_components=k,random_state=seed,learning_method="batch",
                 max_iter=25,evaluate_every=-1,n_jobs=-1,doc_topic_prior=50/k/50,topic_word_prior=.01)
            m.fit(X);ms.append(m);per.append(m.perplexity(X));coh.append(npmi_coherence(m,X));div.append(topic_diversity(m))
        st=stability(ms)
        metrics.append({"k":k,"perplexity_mean":float(np.mean(per)),"perplexity_sd":float(np.std(per)),
          "npmi_mean":float(np.mean(coh)),"npmi_sd":float(np.std(coh)),"diversity_mean":float(np.mean(div)),
          "stability":st})
        all_models[k]=ms
    eligible=[m for m in metrics if m["diversity_mean"]>=.75 and m["stability"]>=.65] or metrics
    best=max(eligible,key=lambda z:(z["npmi_mean"],-z["perplexity_mean"]))
    idx=int(np.argmax([npmi_coherence(m,X) for m in all_models[best["k"]]]))
    return best["k"],all_models[best["k"]][idx],metrics

def topic_rows(model,terms,forms):
    topics=[];align=[]
    for t,w in enumerate(model.components_,1):
        p=w/w.sum(); ids=np.argsort(p)[::-1][:TOP_N]; termset={terms[i] for i in ids}
        bycore=defaultdict(float); matched=set()
        for i in ids:
            term=terms[i]
            for code,lw in forms.get(term,{}).items():
                core=CODE_TO_CORE.get(code)
                if core: bycore[core]+=float(p[i])*lw;matched.add(term)
                align.append([t,term,float(p[i]),code,core or ENABLERS.get(code,"Fuera del modelo nuclear"),lw])
        for cap in CORE:
            capset={f for f,codes in forms.items() if any(CODE_TO_CORE.get(c)==cap for c in codes)}
            inter=termset&capset; union=termset|capset
            j=len(inter)/len(union) if union else 0
            topics.append([t,cap,j,bycore.get(cap,0),"; ".join(sorted(inter)),
                           "; ".join(f"{terms[i]}:{p[i]:.5f}" for i in ids)])
    return topics,align

def document_capability(text,forms):
    tokens=text.split(); uni=Counter(tokens); bi=Counter(" ".join(tokens[i:i+2]) for i in range(len(tokens)-1))
    tri=Counter(" ".join(tokens[i:i+3]) for i in range(len(tokens)-2)); grams={**uni,**bi,**tri}
    has_context=any(c in text for c in CONTEXT); codes=defaultdict(float); evidence=defaultdict(list)
    for form,maps in forms.items():
        n=grams.get(form,0)
        if not n or (form in GENERIC and not has_context):continue
        specificity=1+.20*(len(form.split())-1)
        for code,w in maps.items():
            val=min(n,3)*w*specificity;codes[code]+=val
            if len(evidence[code])<8:evidence[code].append(form)
    core=defaultdict(float)
    for code,val in codes.items():
        if code in CODE_TO_CORE:core[CODE_TO_CORE[code]]+=val
    total=sum(core.values())
    dist={c:(core[c]/total if total else 0) for c in CORE}
    return codes,dist,evidence,total

def profiles(actors,docs,forms):
    byactor=defaultdict(list); docrows=[]
    for d in docs:
        codes,dist,ev,total=document_capability(d["text"],forms)
        byactor[d["actor"]].append((dist,total,codes))
        docrows.append([d["id"],d["actor"],d["year"],total,*[dist[c] for c in CORE],
                        "; ".join(f"{k}:{v:.2f}" for k,v in sorted(codes.items()))])
    signaled=[x for xs in byactor.values() for x in xs if x[1]>0]
    prior={c:(np.mean([x[0][c] for x in signaled]) if signaled else .25) for c in CORE}
    rows=[];oprows=[];vectors={}
    for a in actors:
        xs=byactor[a]; ys=[x for x in xs if x[1]>0]; n=len(ys); N=len(xs)
        raw={c:(np.mean([x[0][c] for x in ys]) if ys else 0) for c in CORE}
        lam=n/(n+10); shr={c:lam*raw[c]+(1-lam)*prior[c] if n else 0 for c in CORE}
        s=sum(shr.values()); shr={c:(shr[c]/s if s else 0) for c in CORE};vectors[a]=shr
        dominant=max(CORE,key=lambda c:shr[c]) if s else "Sin evidencia"
        rows.append([a,N,n,n/N if N else 0,*[raw[c] for c in CORE],*[shr[c] for c in CORE],dominant,
                     "Perfil documental estimado" if n>=10 else ("Evidencia limitada" if n else "Sin evidencia léxica")])
        opdocs=[]
        for _,_,codes in xs:
            z=sum(codes.values())
            if z:opdocs.append({c:codes.get(c,0)/z for c in CODE_ORDER})
        opmean={c:(float(np.mean([q[c] for q in opdocs])) if opdocs else 0) for c in CODE_ORDER}
        opdom=max(CODE_ORDER,key=lambda c:opmean[c]) if opdocs else "Sin evidencia"
        oprows.append([a,N,len(opdocs),*[opmean[c] for c in CODE_ORDER],opdom])
    pairs=[]
    for a,b in itertools.combinations(actors,2):
        x=np.array([vectors[a][c] for c in CORE]);y=np.array([vectors[b][c] for c in CORE])
        den=np.maximum(x,y).sum()
        wj=float(np.minimum(x,y).sum()/den) if x.sum()>0 and y.sum()>0 and den else None
        shared=[c for c in CORE if x[CORE.index(c)]>0 and y[CORE.index(c)]>0]
        pairs.append([a,b,wj,1-wj if wj is not None else None,"; ".join(shared),
          "Similitud de perfiles; la distancia no demuestra complementariedad" if wj is not None else "No evaluable: uno o ambos actores carecen de perfil"])
    return rows,oprows,pairs,docrows

def write_csv(path,header,rows):
    with path.open("w",newline="",encoding="utf-8-sig") as f:
        w=csv.writer(f);w.writerow(header);w.writerows(rows)

def main(corpus,lexicon,out):
    out.mkdir(parents=True,exist_ok=True)
    actors,docs=load_corpus(corpus);lexrows,forms=load_lexicon(lexicon)
    v,X=vectorize([d["text"] for d in docs]);terms=v.get_feature_names_out()
    k,model,metrics=fit_select(X);theta=model.transform(X)
    topics,align=topic_rows(model,terms,forms)
    prows,oprows,pairs,docrows=profiles(actors,docs,forms)
    topic_top=[]
    for t,w in enumerate(model.components_,1):
        p=w/w.sum();ids=np.argsort(p)[::-1][:TOP_N]
        topic_top.append([t,float(theta[:,t-1].mean()),"; ".join(terms[i] for i in ids),
                          "; ".join(f"{terms[i]}:{p[i]:.6f}" for i in ids)])
    write_csv(out/"seleccion_k.csv",["k","perplexity_media","perplexity_sd","npmi_media","npmi_sd","diversidad","estabilidad"],
      [[m["k"],m["perplexity_mean"],m["perplexity_sd"],m["npmi_mean"],m["npmi_sd"],m["diversity_mean"],m["stability"]] for m in metrics])
    write_csv(out/"topicos_lda.csv",["topico","prevalencia","terminos_top","terminos_pesos"],topic_top)
    write_csv(out/"match_topico_ontotic.csv",["topico","capacidad","jaccard_binario","masa_lexica_ponderada","interseccion","terminos_top"],topics)
    write_csv(out/"trazabilidad_match.csv",["topico","termino","peso_topico","codigo","capacidad_nuclear_o_estado","peso_lexico"],align)
    write_csv(out/"perfiles_actores.csv",["actor","documentos_lda","documentos_con_señal","cobertura_señal",
      *[f"bruto_{c}" for c in CORE],*[f"ajustado_{c}" for c in CORE],"capacidad_predominante","calidad_evidencia"],prows)
    write_csv(out/"perfiles_operativos_12.csv",["actor","documentos_lda","documentos_con_señal",*CODE_ORDER,
      "codigo_predominante"],oprows)
    write_csv(out/"jaccard_actores.csv",["actor_a","actor_b","jaccard_ponderado","distancia_perfil","dimensiones_compartidas","interpretacion"],pairs)
    write_csv(out/"señales_documento.csv",["documento_id","actor","año","puntaje_lexico",*CORE,"codigos_operativos"],docrows)
    meta={"version":"V6.13","documentos":len(docs),"actores":len(actors),"formas_lexicas":len(forms),
      "registros_lexicos":len(lexrows),"features":int(X.shape[1]),"k_seleccionado":k,"semillas":SEEDS,
      "criterio_k":"Máximo NPMI entre modelos con diversidad >=0.75 y estabilidad >=0.65; perplexity como desempate",
      "jaccard_topico":"Jaccard binario entre top-20 del tópico y formas léxicas de la capacidad",
      "jaccard_actores":"Jaccard ponderado=sum(min perfiles)/sum(max perfiles); similitud, no complementariedad",
      "perfil":"Media documental normalizada con contracción empírico-bayesiana n/(n+10)",
      "habilitadoras_no_forzadas":ENABLERS,"mapeo_codigos_nucleares":CODE_TO_CORE,
      "limitaciones":["El match léxico requiere validación experta estratificada","LDA produce temas, no capacidades",
        "Jaccard no acredita relaciones, ajuste funcional ni complementariedad realizada"]}
    (out/"metodologia.json").write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(meta,ensure_ascii=False,indent=2))

if __name__=="__main__":
    main(Path(sys.argv[1]),Path(sys.argv[2]),Path(sys.argv[3]))
